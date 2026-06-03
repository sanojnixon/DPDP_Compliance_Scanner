"""
ropa_export.py — Excel Export Engine
=====================================
Loads the SIB ROPA template, injects populated field values,
and generates a downloadable .xlsx file.
"""

import io
import logging
import os
import copy
from datetime import datetime
from typing import Any, Dict, List, Optional

import openpyxl
from openpyxl.utils import get_column_letter

logger = logging.getLogger(__name__)

# Template lives alongside the backend scripts (copied from ropa/ at build time)
_TEMPLATE_FILE = os.path.join(
    os.path.dirname(os.path.abspath(__file__)),
    "ropa_template.xlsx",
)

# Column key → Excel column index (1-based) in Fiduciary_RoPA sheet
_COL_INDEX = {
    "serial": 1, "department": 2, "spoc_name": 3, "application_name": 4,
    "process": 5, "purpose": 6, "description": 7, "data_principal_cats": 8,
    "includes_children": 9, "child_tracking": 10, "personal_data_cats": 11,
    "data_source": 12, "grounds": 13, "consent_type": 14,
    "automated_decisions": 15, "data_principal_rights": 16,
    "processor_name": 17, "dpa_in_place": 18, "third_party_safeguards": 19,
    "other_recipients": 20, "transfer_outside": 21, "transfer_country": 22,
    "transfer_purpose": 23, "transfer_safeguards": 24,
    "storage_format": 25, "storage_location": 26,
    "retention_period": 27, "retention_purpose": 28,
    "dpia_applicable": 29, "dpia_link": 30, "privacy_notice": 31,
}

# Personal Data Inventory — row mapping in template (sheet index 3)
# Category → start_row in the template
_INVENTORY_ROWS = {
    "KYC Documents": {
        "PAN Details": 11,
        "Aadhar details": 12,
        "Passport details": 13,
        "Voter ID": 14,
        "Driving licence details": 15,
    },
    "Identity Data": {
        "Name": 18,
        "Identification number (including Customer ID/Employee ID)": 19,
        "Personal Email ID": 31,
        "Personal Phone Number": 32,
    },
    "Biometric Data": {
        "Fingerprints": 55,
        "Voice patterns": 56,
    },
    "Online Identifiers": {
        "Username": 57,
        "Passwords": 58,
        "IP Address": 59,
        "Cookie details (location, browsing patterns, etc.)": 60,
    },
    "Financial Data": {
        "Bank account numbers": 69,
        "Debit Card details": 70,
        "Credit Card details": 71,
        "UPI Handles": 74,
        "Account balance": 75,
        "Account transaction history": 76,
    },
    "Children and Person with Disability Data": {
        "Child Name": 61,
        "Child Date of Birth": 62,
    },
}


def _load_template() -> openpyxl.Workbook:
    """Load the ROPA template workbook."""
    if not os.path.exists(_TEMPLATE_FILE):
        raise FileNotFoundError(f"ROPA template not found at {_TEMPLATE_FILE}")
    wb = openpyxl.load_workbook(_TEMPLATE_FILE)
    logger.info("[ROPA Export] Template loaded: %s", _TEMPLATE_FILE)
    return wb


def export_ropa_xlsx(
    ropa_entries: List[Dict[str, Any]],
    personal_data_inventory: Optional[Dict[str, Any]] = None,
    metadata: Optional[Dict[str, Any]] = None,
) -> io.BytesIO:
    """
    Generate a populated ROPA Excel file from entries.

    Args:
        ropa_entries: List of ROPA entry dicts (one per row).
        personal_data_inventory: Optional PII inventory data.
        metadata: Optional metadata (date, version, etc.).

    Returns:
        BytesIO buffer containing the .xlsx file.
    """
    metadata = metadata or {}
    wb = _load_template()

    # ── 1. Populate Fiduciary_RoPA sheet ──────────────────────────────────
    ws = wb["Fiduciary_RoPA"]
    data_start_row = 3  # Row 1-2 are headers

    for i, entry in enumerate(ropa_entries):
        row = data_start_row + i
        for key, col_idx in _COL_INDEX.items():
            value = entry.get(key, "")
            if value is not None and value != "":
                cell = ws.cell(row=row, column=col_idx)
                cell.value = str(value)
                # Apply text wrapping for long content
                cell.alignment = openpyxl.styles.Alignment(
                    wrap_text=True, vertical="top"
                )

    logger.info(
        "[ROPA Export] Populated %d rows in Fiduciary_RoPA sheet",
        len(ropa_entries),
    )

    # ── 2. Populate Personal Data Inventory sheet ─────────────────────────
    if personal_data_inventory:
        _populate_inventory(wb, personal_data_inventory)

    # ── 3. Update Coverpage metadata ──────────────────────────────────────
    if "Coverpage" in wb.sheetnames:
        cover = wb["Coverpage"]
        # Date field (row 7)
        cover.cell(row=7, column=2).value = metadata.get(
            "date", datetime.now().strftime("%d-%b-%Y")
        )
        # Version (row 8)
        if metadata.get("version"):
            cover.cell(row=8, column=2).value = metadata["version"]
        # Department (row 3)
        if ropa_entries and ropa_entries[0].get("department"):
            cover.cell(row=3, column=2).value = ropa_entries[0]["department"]
        # SPOC (row 4)
        if ropa_entries and ropa_entries[0].get("spoc_name"):
            cover.cell(row=4, column=2).value = ropa_entries[0]["spoc_name"]

    # ── 4. Write to buffer ────────────────────────────────────────────────
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    wb.close()

    logger.info("[ROPA Export] Excel file generated successfully (%d bytes)", buffer.getbuffer().nbytes)
    return buffer


def _populate_inventory(wb: openpyxl.Workbook, inventory: Dict[str, Any]):
    """Populate the New Personal Data Inventory sheet."""
    if "New Personal Data Inventory" not in wb.sheetnames:
        logger.warning("[ROPA Export] 'New Personal Data Inventory' sheet not found")
        return

    ws = wb["New Personal Data Inventory"]

    for category, items in inventory.items():
        if category not in _INVENTORY_ROWS:
            continue
        row_map = _INVENTORY_ROWS[category]
        for item in items:
            param = item.get("parameter", "")
            if param in row_map:
                row = row_map[param]
                # Column C = Department Response (Yes/No)
                ws.cell(row=row, column=3).value = item.get("response", "Yes")
                # Column D = Department Remarks
                ws.cell(row=row, column=4).value = item.get("remarks", "")
                # Column E = Mode of collection
                if ws.max_column >= 5:
                    ws.cell(row=row, column=5).value = item.get("mode_collection", "Digital")
                # Column F = Mode of Processing
                if ws.max_column >= 6:
                    ws.cell(row=row, column=6).value = item.get("mode_processing", "Digital")

    logger.info("[ROPA Export] Personal Data Inventory populated")
